"""
报告自动化（对应 think-cell Advanced report automation）：
一个 JSON（或 Python dict）描述整份报告 → 一条命令批量输出 HTML / PPTX / XLSX / SVG / PNG。
数据更新时只改 JSON 或 CSV，重跑即可（替代 think-cell 的 Excel 数据链接）。

用法：
    python -m tccore.report spec.json 输出目录 [--formats html,pptx,xlsx,svg,png]

spec 结构见 SKILL.md「报告 JSON 规格」。
"""
from __future__ import annotations

import csv
import json
import math
import os
import sys

from . import annotations as A
from . import charts_bar, charts_gantt, charts_line, charts_mekko, charts_round, charts_xy
from .table import table as _table_fn
from .figure import Figure, html_page

CHARTS = {
    "column": charts_bar.column, "bar": lambda f, **k: charts_bar.column(f, horizontal=True, **k),
    "butterfly": charts_bar.butterfly, "waterfall": charts_bar.waterfall,
    "line": charts_line.line, "area": charts_line.area, "combo": charts_line.combo, "pareto": charts_line.pareto,
    "candlestick": charts_line.candlestick, "football": charts_line.football,
    "pie": charts_round.pie, "pie_of_pie": charts_round.pie_of_pie, "concentric": charts_round.concentric,
    "gauge": charts_round.gauge, "scatter": charts_xy.scatter, "bubble": lambda f, **k: charts_xy.scatter(f, bubble=True, **k),
    "mekko": charts_mekko.mekko, "gantt": charts_gantt.gantt, "table": _table_fn,
}
ANN = {"cagr": A.cagr_arrow, "diff": A.diff_arrow, "value_line": A.value_line, "series_cagr": A.series_cagr,
       "cagr_column": A.cagr_column, "series_connectors": A.series_connectors, "callout": A.callout}


def load_csv(path, cat_col=0):
    """CSV → (cats, {系列: [值]})。第一列为类目，其余列为系列；空单元格 → None。"""
    with open(path, encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    head, body = rows[0], rows[1:]
    cats = [r[cat_col] for r in body]
    series = {}
    for j, h in enumerate(head):
        if j == cat_col:
            continue
        vals = []
        for r in body:
            s = r[j].strip().replace(",", "") if j < len(r) else ""
            vals.append(None if s == "" else ("e" if s == "e" else float(s)))
        series[h] = vals
    return cats, series


def _tuplify(v):
    if isinstance(v, list) and len(v) == 2 and all(isinstance(x, int) for x in v):
        return tuple(v)
    if isinstance(v, list) and len(v) == 2 and v[0] == "v":
        return tuple(v)
    return v


def build_figure(ch: dict, theme=None, base_dir="."):
    def validate_finite(value, path="chart"):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError(f"{path} contains a non-finite number")
        if isinstance(value, list):
            for i, item in enumerate(value):
                validate_finite(item, f"{path}[{i}]")
        elif isinstance(value, dict):
            for key, item in value.items():
                validate_finite(item, f"{path}.{key}")

    validate_finite(ch)
    ch = dict(ch)
    W, H = ch.pop("size", [680, 400])
    fig = Figure(W, H, ch.pop("title", None), ch.pop("subtitle", None), ch.pop("source", None),
                 theme=ch.pop("theme", theme))
    kind = ch.pop("type")
    if kind not in CHARTS:
        raise ValueError(f"unsupported chart type: {kind}")
    anns = ch.pop("annotations", [])
    ch.pop("note", None)
    ch.pop("id", None)
    pmode = ch.pop("pptx", "shapes")
    if "csv" in ch:
        cats, series = load_csv(os.path.join(base_dir, ch.pop("csv")))
        if kind == "waterfall":
            ch.setdefault("steps", list(zip(cats, list(series.values())[0])))
        else:
            ch.setdefault("cats", cats)
            ch.setdefault("series", series)
    if kind == "waterfall" and "steps" in ch:
        ch["steps"] = [tuple(s) for s in ch["steps"]]
    for k in ("left", "right"):
        if k in ch and isinstance(ch[k], list):
            ch[k] = tuple(ch[k])
    if "highlight" in ch and isinstance(ch["highlight"], dict):
        hl = {}
        for k, v in ch["highlight"].items():
            key = tuple(int(x) for x in k.split(",")) if "," in str(k) else (int(k) if str(k).isdigit() else k)
            hl[key] = fig.t["ACCENT"] if v in ("accent", True) else v
        ch["highlight"] = hl
    if "gaps" in ch:
        ch["gaps"] = {int(k): v for k, v in ch["gaps"].items()}
    if "breaks" in ch:
        ch["breaks"] = [tuple(b) for b in ch["breaks"]]
    if "rows" in ch and kind == "football":
        ch["rows"] = [tuple(r) for r in ch["rows"]]
    if "ohlc" in ch:
        ch["ohlc"] = [tuple(r) for r in ch["ohlc"]]
    ctx = CHARTS[kind](fig, **ch)
    for a in anns:
        a = dict(a)
        typ = a.pop("type")
        fn = ANN[typ]
        for k in list(a):
            a[k] = _tuplify(a[k])
        res = fn(ctx, **a)
        if res is not None:
            fig.meta.append({"annotation": typ, "args": {k: v for k, v in a.items()}, "value": res})
    return fig


# ---------------------------------------------------------------- 各载体的原生数据
def _native_payload(ch, base_dir="."):
    """返回 (xlsx/pptx 原生图 kind, cats, series, extra) 或 None（只能贴图）。"""
    kind = ch["type"]
    d = dict(ch)
    if "csv" in d:
        cats, series = load_csv(os.path.join(base_dir, d["csv"]))
        d.setdefault("cats", cats)
        d.setdefault("series", series)
        if kind == "waterfall":
            d.setdefault("steps", list(zip(cats, list(series.values())[0])))
    if kind in ("column", "bar"):
        mode = d.get("mode", "stacked")
        hz = kind == "bar" or d.get("horizontal")
        base = "bar" if hz else "column"
        k = base + {"stacked": "_stacked", "clustered": "", "100": "_100"}[mode]
        if len(d["series"]) == 1 and mode == "stacked":
            k = base
        return k, d["cats"], d["series"], {}
    if kind == "line":
        return "line", d["cats"], d["series"], {}
    if kind == "area":
        return {"stacked": "area_stacked", "100": "area_100", "overlap": "area"}[d.get("mode", "stacked")], d["cats"], d["series"], {}
    if kind == "pie":
        return ("doughnut" if d.get("donut") else "pie"), d["labels"], {"值": d["values"]}, {}
    if kind == "waterfall":
        steps = d["steps"]
        return "waterfall", [s[0] for s in steps], {"值": [s[1] if not isinstance(s[1], list) else sum(s[1]) for s in steps]}, {}
    if kind == "combo":
        s = dict(d["bars"])
        s.update(d["lines"])
        return "combo", d["cats"], s, {"secondary": list(d["lines"])}
    if kind in ("scatter", "bubble"):
        pts = d["points"]
        s = {"x": [p["x"] for p in pts], "y": [p["y"] for p in pts]}
        if kind == "bubble":
            s["size"] = [p.get("s", 1) for p in pts]
        return kind, [p.get("label") or "" for p in pts], s, {}
    if kind == "gantt":
        rows = [r for r in d["rows"] if r.get("start") and r.get("end")]
        return "gantt", [r["name"] for r in rows], {"start": [r["start"] for r in rows], "end": [r["end"] for r in rows]}, {}
    return None


def _data_table(ch, base_dir="."):
    """只能贴图的类型：把源数据整理成表格写进 Excel。"""
    kind, d = ch["type"], ch
    if "csv" in d:
        cats, series = load_csv(os.path.join(base_dir, d["csv"]))
        return cats, series
    if kind == "mekko":
        return d["cats"], d["series"]
    if kind == "butterfly":
        return d["cats"], {d["left"][0]: d["left"][1], d["right"][0]: d["right"][1]}
    if kind == "candlestick":
        o = d["ohlc"]
        return d["cats"], {"开": [r[0] for r in o], "高": [r[1] for r in o], "低": [r[2] for r in o], "收": [r[3] for r in o]}
    if kind == "football":
        return [r[0] for r in d["rows"]], {"低": [r[1] for r in d["rows"]], "高": [r[2] for r in d["rows"]]}
    if kind in ("pie_of_pie",):
        return d["labels"], {"值": d["values"]}
    if kind == "concentric":
        return d["labels"], d["series"]
    if kind == "pareto":
        return d["cats"], {"值": d["values"]}
    if kind == "gauge":
        return ["值"], {"值": [d["value"]]}
    if kind == "table":
        return [r[0] for r in d["rows"]], {h: [r[j + 1] for r in d["rows"]] for j, h in enumerate(d["header"][1:])}
    return None


# ---------------------------------------------------------------- 构建
def build(spec, out_dir, formats=("html", "pptx", "xlsx"), base_dir="."):
    if isinstance(spec, str):
        base_dir = os.path.dirname(os.path.abspath(spec))
        spec = json.load(open(spec, encoding="utf-8"))
    os.makedirs(out_dir, exist_ok=True)
    theme = spec.get("theme")
    name = spec.get("file", "report")
    charts = spec["charts"]
    figs = [build_figure(ch, theme, base_dir) for ch in charts]
    out = {}
    if "html" in formats:
        p = os.path.join(out_dir, f"{name}.html")
        open(p, "w", encoding="utf-8").write(
            html_page(figs, spec.get("title", "报告"), {i: ch["note"] for i, ch in enumerate(charts) if ch.get("note")},
                      cols=spec.get("columns", 1)))
        out["html"] = p
    if "svg" in formats or "png" in formats:
        for i, (ch, f) in enumerate(zip(charts, figs)):
            cid = ch.get("id", f"chart{i + 1:02d}")
            if "svg" in formats:
                f.save_svg(os.path.join(out_dir, f"{cid}.svg"))
            if "png" in formats:
                f.png(os.path.join(out_dir, f"{cid}.png"))
        out["svg/png"] = out_dir
    if "pptx" in formats:
        out["pptx"] = _build_pptx(spec, charts, figs, os.path.join(out_dir, f"{name}.pptx"), base_dir)
    if "xlsx" in formats:
        out["xlsx"] = _build_xlsx(spec, charts, figs, os.path.join(out_dir, f"{name}.xlsx"), base_dir)
    return out


def _build_pptx(spec, charts, figs, path, base_dir):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from .pptx_native import native_chart
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    t = figs[0].t if figs else None
    if spec.get("title"):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(11.5), Inches(1.2))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = spec["title"]
        r.font.size, r.font.bold = Pt(32), True
        if spec.get("subtitle"):
            p2 = tb.text_frame.add_paragraph()
            r2 = p2.add_run()
            r2.text = spec["subtitle"]
            r2.font.size = Pt(16)
    for ch, f in zip(charts, figs):
        mode = ch.get("pptx", "shapes")
        if mode in ("shapes", "both"):
            s = prs.slides.add_slide(blank)
            wmax, hmax = 12.3, 6.9
            w_in = min(wmax, hmax * f.W / f.H)
            f.pptx(s, (13.333 - w_in) / 2, 0.3, w_in=w_in)
            if ch.get("note"):
                s.notes_slide.notes_text_frame.text = ch["note"]
        if mode in ("native", "both"):
            nat = _native_payload(ch, base_dir)
            if nat and nat[0] not in ("waterfall", "combo", "gantt"):
                s = prs.slides.add_slide(blank)
                tb = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12), Inches(0.8))
                r = tb.text_frame.paragraphs[0].add_run()
                r.text = ch.get("title", "")
                r.font.size, r.font.bold = Pt(20), True
                kind, cats, series, extra = nat
                if kind in ("scatter", "bubble"):
                    pts = list(zip(series["x"], series["y"], series.get("size", [1] * len(series["x"]))))
                    native_chart(s, kind, points=[p if kind == "bubble" else p[:2] for p in pts], theme=t)
                else:
                    native_chart(s, kind, cats, series, theme=t)
    prs.save(path)
    return path


def _build_xlsx(spec, charts, figs, path, base_dir):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from .xlsx import add_figure, native_chart, write_table
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for i, (ch, f) in enumerate(zip(charts, figs)):
        nm = (ch.get("id") or f"图{i + 1}")[:28]
        while nm in used:
            nm += "_"
        used.add(nm)
        ws = wb.create_sheet(nm)
        ws["A1"] = ch.get("title", "")
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = ch.get("subtitle", "")
        ws["A2"].font = Font(color="8C8C8C")
        nat = _native_payload(ch, base_dir)
        row_img = 4
        if nat:
            kind, cats, series, extra = nat
            native_chart(ws, kind, cats, series, anchor="H4", title="原生图表（可改数据）", t=f.t, start_row=4,
                         secondary=extra.get("secondary"))
            row_img = max(26, len(cats) + 8)
        else:
            dt = _data_table(ch, base_dir)
            if dt:
                write_table(ws, dt[0], dt[1], 4, 1, t=f.t)
                row_img = len(dt[0]) + 8
        ws.cell(row_img - 1, 1, "高保真图（含全部 think-cell 注释）").font = Font(bold=True, color="8C8C8C")
        add_figure(ws, f, f"A{row_img}")
    wb.save(path)
    return path


def main(argv=None):
    argv = argv or sys.argv[1:]
    if len(argv) < 2:
        print(__doc__)
        return 1
    fm = ("html", "pptx", "xlsx")
    if "--formats" in argv:
        fm = tuple(argv[argv.index("--formats") + 1].split(","))
    res = build(argv[0], argv[1], fm)
    for k, v in res.items():
        print(f"{k:8s} {v}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
