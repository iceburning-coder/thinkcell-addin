"""
Excel 后端（openpyxl）：
  write_table     源数据写入工作表（带表头样式）
  native_chart    原生图表（用户可改数据）：column/bar(stacked/clustered/100)、line、area、pie、doughnut、
                  scatter、bubble、combo（柱+线次轴）、waterfall（透明底座）、gantt（堆积条+日期）
  add_figure      高保真图片（Figure → PNG），带全部 think-cell 注释
Excel 原生图表做不了的（CAGR/差异箭头、数值线、Mekko、断轴…）一律用 add_figure 贴高保真图。
"""
from __future__ import annotations

import os
import tempfile
from datetime import date

from openpyxl.chart import (AreaChart, BarChart, BubbleChart, DoughnutChart, LineChart, PieChart, Reference,
                            ScatterChart, Series)
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .style import series_colors, theme as _theme


def _hex(c):
    return c.lstrip("#").upper()


def write_table(ws, cats, series: dict, start_row=1, start_col=1, cat_header="", number_format="#,##0.0", t=None):
    t = t or _theme()
    hdr = [cat_header] + list(series)
    for j, h in enumerate(hdr):
        cell = ws.cell(start_row, start_col + j, h)
        cell.font = Font(bold=True, color="FFFFFF", name="Arial")
        cell.fill = PatternFill("solid", fgColor=_hex(t["SERIES"][0]))
        cell.alignment = Alignment(horizontal="center")
    for i, cat in enumerate(cats):
        ws.cell(start_row + 1 + i, start_col, cat).font = Font(name="Arial")
        for j, nm in enumerate(series):
            v = series[nm][i]
            cell = ws.cell(start_row + 1 + i, start_col + 1 + j, v)
            cell.number_format = number_format
            cell.font = Font(name="Arial")
            if i % 2 == 1:
                cell.fill = PatternFill("solid", fgColor="F4F4F4")
    ws.column_dimensions[get_column_letter(start_col)].width = max(10, max(len(str(c)) for c in cats) * 1.8)
    for j in range(len(series)):
        ws.column_dimensions[get_column_letter(start_col + 1 + j)].width = 12
    return (start_row, start_col, start_row + len(cats), start_col + len(series))


def _style_chart(ch, t, title=None, legend=True, w_cm=17, h_cm=9.5):
    ch.title = title
    ch.width, ch.height = w_cm, h_cm
    ch.style = 2
    if not legend:
        ch.legend = None
    else:
        ch.legend.position = "t"
    return ch


def _no_grid(ch):
    try:
        ch.y_axis.majorGridlines = None
    except Exception:
        pass
    ch.y_axis.delete = False
    ch.x_axis.delete = False


def native_chart(ws, kind, cats, series: dict, anchor="H2", title=None, t=None, start_row=1, start_col=1,
                 labels=True, number_format="#,##0.0", secondary=None, dates=None):
    """
    kind: column | column_stacked | column_100 | bar | bar_stacked | bar_100 | line | area | area_stacked |
          area_100 | pie | doughnut | scatter | bubble | combo | waterfall | gantt
    series: {系列: [值]}；combo 时 secondary=[折线系列名]
    """
    t = t or _theme()
    if kind == "waterfall":
        return _waterfall(ws, cats, series, anchor, title, t, start_row, start_col)
    if kind == "gantt":
        return _gantt(ws, cats, series, anchor, title, t, start_row, start_col)
    r0, c0, r1, c1 = write_table(ws, cats, series, start_row, start_col, number_format=number_format, t=t)
    catref = Reference(ws, min_col=c0, min_row=r0 + 1, max_row=r1)
    names = list(series)
    cols = series_colors(t, len(names))
    if kind.startswith(("column", "bar")):
        ch = BarChart()
        ch.type = "col" if kind.startswith("column") else "bar"
        ch.grouping = "percentStacked" if kind.endswith("100") else ("stacked" if kind.endswith("stacked") else "clustered")
        if ch.grouping != "clustered":
            ch.overlap = 100
        ch.gapWidth = 60
    elif kind.startswith("area"):
        ch = AreaChart()
        ch.grouping = "percentStacked" if kind.endswith("100") else ("stacked" if kind.endswith("stacked") else "standard")
    elif kind == "line":
        ch = LineChart()
    elif kind in ("pie", "doughnut"):
        ch = PieChart() if kind == "pie" else DoughnutChart(holeSize=55)
    elif kind in ("scatter", "bubble"):
        ch = ScatterChart() if kind == "scatter" else BubbleChart()
    elif kind == "combo":
        ch = BarChart()
        ch.grouping, ch.gapWidth = "clustered", 60
    else:
        raise ValueError(kind)
    sec = set(secondary or [])
    if kind in ("scatter", "bubble"):
        # 约定：series = {"x": [...], "y": [...], "size": [...]}
        xs = Reference(ws, min_col=c0 + 1, min_row=r0 + 1, max_row=r1)
        ys = Reference(ws, min_col=c0 + 2, min_row=r0 + 1, max_row=r1)
        if kind == "bubble":
            sz = Reference(ws, min_col=c0 + 3, min_row=r0 + 1, max_row=r1)
            s = Series(values=ys, xvalues=xs, zvalues=sz, title="")
        else:
            s = Series(ys, xs, title="")
            s.marker.symbol = "circle"
            s.graphicalProperties.line.noFill = True
        s.graphicalProperties.solidFill = _hex(cols[0])
        ch.series.append(s)
    else:
        main_names = [n for n in names if n not in sec]
        for j, nm in enumerate(names):
            if nm in sec:
                continue
            data = Reference(ws, min_col=c0 + 1 + j, min_row=r0, max_row=r1)
            ch.add_data(data, titles_from_data=True)
        ch.set_categories(catref)
        for j, s in enumerate(ch.series):
            col = _hex(cols[j % len(cols)])
            if kind == "line":
                s.graphicalProperties.line.solidFill = col
                s.graphicalProperties.line.width = 28000
                s.smooth = False
            elif kind not in ("pie", "doughnut"):
                s.graphicalProperties.solidFill = col
                s.graphicalProperties.line.solidFill = "FFFFFF"
        if kind in ("pie", "doughnut"):
            from openpyxl.chart.series import DataPoint
            pc = series_colors(t, len(cats))
            for i in range(len(cats)):
                pt = DataPoint(idx=i)
                pt.graphicalProperties.solidFill = _hex(pc[i])
                pt.graphicalProperties.line.solidFill = "FFFFFF"
                ch.series[0].dPt.append(pt)
        if labels:
            ch.dataLabels = DataLabelList()
            ch.dataLabels.showVal = kind not in ("pie", "doughnut")
            ch.dataLabels.showPercent = kind in ("pie", "doughnut")
            ch.dataLabels.numFmt = number_format
            ch.dataLabels.showSerName = ch.dataLabels.showCatName = ch.dataLabels.showLegendKey = False
        if kind == "combo" and sec:
            lc = LineChart()
            for j, nm in enumerate(names):
                if nm in sec:
                    lc.add_data(Reference(ws, min_col=c0 + 1 + j, min_row=r0, max_row=r1), titles_from_data=True)
            for s in lc.series:
                s.graphicalProperties.line.solidFill = _hex(t["ACCENT"])
                s.graphicalProperties.line.width = 28000
                s.marker.symbol = "circle"
            lc.y_axis.axId = 200
            lc.y_axis.crosses = "max"
            lc.y_axis.majorGridlines = None
            ch += lc
    _style_chart(ch, t, title, legend=len(names) > 1 or kind in ("pie", "doughnut"))
    if kind not in ("pie", "doughnut"):
        _no_grid(ch)
    ws.add_chart(ch, anchor)
    return ch


def _waterfall(ws, cats, series, anchor, title, t, r0, c0):
    """瀑布：透明底座 + 增/减/合计三列堆积。series = {"值": [数值或 'e']}"""
    vals = list(series.values())[0]
    base, up, down, tot = [], [], [], []
    run = 0.0
    for i, v in enumerate(vals):
        if i == 0 and v != "e":
            run = float(v)
            base.append(0); up.append(None); down.append(None); tot.append(run)
        elif v == "e":
            base.append(0); up.append(None); down.append(None); tot.append(run)
        else:
            v = float(v)
            if v >= 0:
                base.append(run); up.append(v); down.append(None)
            else:
                base.append(run + v); up.append(None); down.append(-v)
            tot.append(None)
            run += v
    data = {"底座": base, "增加": up, "减少": down, "合计": tot}
    rr0, cc0, rr1, cc1 = write_table(ws, cats, data, r0, c0, t=t)
    ws.cell(rr0, cc1 + 2, "原始数据").font = Font(bold=True)
    for i, v in enumerate(vals):
        ws.cell(rr0 + 1 + i, cc1 + 2, v)
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap, ch.gapWidth = "col", "stacked", 100, 50
    ch.add_data(Reference(ws, min_col=cc0 + 1, max_col=cc1, min_row=rr0, max_row=rr1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cc0, min_row=rr0 + 1, max_row=rr1))
    colors = [None, t["POS"], t["NEG"], t["TOTAL"]]
    for s, col in zip(ch.series, colors):
        if col is None:
            s.graphicalProperties.noFill = True
            s.graphicalProperties.line.noFill = True
        else:
            s.graphicalProperties.solidFill = _hex(col)
            s.dLbls = DataLabelList()
            s.dLbls.showVal = True
            s.dLbls.numFmt = "#,##0.0"
            s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    _style_chart(ch, t, title, legend=False)
    _no_grid(ch)
    ws.add_chart(ch, anchor)
    return ch


def _gantt(ws, cats, series, anchor, title, t, r0, c0):
    """甘特：series = {"start": [date], "end": [date]}，透明起始 + 工期条，横轴为日期。"""
    starts = series["start"]
    ends = series["end"]
    from .axes import as_date
    s_d = [as_date(d) for d in starts]
    e_d = [as_date(d) for d in ends]
    data = {"开始": s_d, "工期(天)": [(b - a).days for a, b in zip(s_d, e_d)]}
    rr0, cc0, rr1, cc1 = write_table(ws, cats, data, r0, c0, number_format="General", t=t)
    for i in range(len(cats)):
        ws.cell(rr0 + 1 + i, cc0 + 1).number_format = "yyyy-mm-dd"
        ws.cell(rr0 + 1 + i, cc0 + 2).number_format = "0"
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap, ch.gapWidth = "bar", "stacked", 100, 40
    ch.add_data(Reference(ws, min_col=cc0 + 1, max_col=cc0 + 2, min_row=rr0, max_row=rr1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cc0, min_row=rr0 + 1, max_row=rr1))
    ch.series[0].graphicalProperties.noFill = True
    ch.series[0].graphicalProperties.line.noFill = True
    ch.series[1].graphicalProperties.solidFill = _hex(t["SERIES"][1])
    ch.x_axis.scaling.orientation = "maxMin"
    ch.y_axis.number_format = "yy/mm"
    ch.y_axis.scaling.min = (min(s_d) - date(1899, 12, 30)).days
    ch.y_axis.scaling.max = (max(e_d) - date(1899, 12, 30)).days
    _style_chart(ch, t, title, legend=False)
    ws.add_chart(ch, anchor)
    return ch


def add_figure(ws, fig, anchor="A20", width_px=760, scale=2.0):
    """贴入高保真 PNG（包含全部注释）。"""
    fd, path = tempfile.mkstemp(suffix=".png")
    os.close(fd)
    fig.png(path, scale=scale)
    img = XLImage(path)
    ratio = width_px / img.width
    img.width, img.height = width_px, int(img.height * ratio)
    ws.add_image(img, anchor)
    return img
