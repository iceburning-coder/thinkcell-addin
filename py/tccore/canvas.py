"""
图元画布 —— 所有图表都只往这里画基本图元，再由后端翻译：
    to_svg()           -> SVG 字符串（HTML/Obsidian）
    to_png(path)       -> PNG（matplotlib 光栅化）
    to_pptx(slide,...) -> PPT 原生形状（组合成一个组，可编辑）
坐标单位 pt，原点左上角，y 向下。
"""
from __future__ import annotations

import math
import html
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Optional


# ---------------------------------------------------------------- 文字宽度估算
def char_w(ch: str) -> float:
    if unicodedata.east_asian_width(ch) in ("W", "F"):
        return 1.0
    if ch in "il.,:;'|!":
        return 0.28
    if ch in " ()[]-/":
        return 0.36
    if ch in "mwMW%":
        return 0.85
    if ch.isupper():
        return 0.66
    return 0.56


def text_w(s: str, size: float, bold: bool = False) -> float:
    w = sum(char_w(c) for c in str(s)) * size
    return w * (1.05 if bold else 1.0)


@dataclass
class Prim:
    kind: str
    a: dict = field(default_factory=dict)


class Canvas:
    def __init__(self, W: float, H: float, theme: dict):
        self.W, self.H, self.t = W, H, theme
        self.items: list[Prim] = []

    # ---------------- 基本图元
    def rect(self, x, y, w, h, fill=None, stroke=None, sw=0.0, r=0.0, dash=None, opacity=1.0):
        if w < 0:
            x, w = x + w, -w
        if h < 0:
            y, h = y + h, -h
        self.items.append(Prim("rect", dict(x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, sw=sw,
                                            r=r, dash=dash, opacity=opacity)))

    def line(self, x1, y1, x2, y2, color, sw=0.75, dash=None):
        self.items.append(Prim("line", dict(x1=x1, y1=y1, x2=x2, y2=y2, color=color, sw=sw, dash=dash)))

    def poly(self, pts, fill=None, stroke=None, sw=0.0, closed=True, dash=None, opacity=1.0):
        self.items.append(Prim("poly", dict(pts=[(float(a), float(b)) for a, b in pts], fill=fill,
                                            stroke=stroke, sw=sw, closed=closed, dash=dash,
                                            opacity=opacity)))

    def circle(self, cx, cy, r, fill=None, stroke=None, sw=0.0, opacity=1.0):
        self.items.append(Prim("circle", dict(cx=cx, cy=cy, r=r, fill=fill, stroke=stroke, sw=sw,
                                              opacity=opacity)))

    def text(self, x, y, s, size=None, color=None, anchor="middle", valign="middle", bold=False,
             rotate=0, box=None, italic=False):
        """
        anchor: start|middle|end（水平），valign: top|middle|bottom（垂直）
        box: dict(fill, stroke, r, pad) 给文字加底框（CAGR/差异标签用）
        """
        s = str(s)
        if s == "":
            return
        size = size or self.t["SIZE_LABEL"]
        color = color or self.t["FG"]
        self.items.append(Prim("text", dict(x=x, y=y, s=s, size=size, color=color, anchor=anchor,
                                            valign=valign, bold=bold, rotate=rotate, box=box,
                                            italic=italic)))

    # ---------------- 复合图元
    def wedge(self, cx, cy, r_out, a0, a1, fill, r_in=0.0, stroke="#FFFFFF", sw=0.8, n=None):
        """扇形/环形段。角度：度，0 = 12 点方向，顺时针。"""
        n = n or max(6, int(abs(a1 - a0) / 3))
        def pt(r, a):
            rad = math.radians(a - 90)
            return (cx + r * math.cos(rad), cy + r * math.sin(rad))
        outer = [pt(r_out, a0 + (a1 - a0) * i / n) for i in range(n + 1)]
        inner = [pt(r_in, a1 - (a1 - a0) * i / n) for i in range(n + 1)] if r_in > 0 else [(cx, cy)]
        self.poly(outer + inner, fill=fill, stroke=stroke, sw=sw)

    def arrow(self, x1, y1, x2, y2, color=None, sw=0.9, head="end", hs=5.0, dash=None):
        """带箭头的线。head: end|start|both|none"""
        color = color or self.t["FG"]
        ang = math.atan2(y2 - y1, x2 - x1)
        def tri(xe, ye, a):
            p1 = (xe, ye)
            p2 = (xe - hs * math.cos(a) + hs * 0.45 * math.sin(a), ye - hs * math.sin(a) - hs * 0.45 * math.cos(a))
            p3 = (xe - hs * math.cos(a) - hs * 0.45 * math.sin(a), ye - hs * math.sin(a) + hs * 0.45 * math.cos(a))
            self.poly([p1, p2, p3], fill=color)
        sx, sy, ex, ey = x1, y1, x2, y2
        if head in ("end", "both"):
            ex, ey = x2 - hs * 0.9 * math.cos(ang), y2 - hs * 0.9 * math.sin(ang)
        if head in ("start", "both"):
            sx, sy = x1 + hs * 0.9 * math.cos(ang), y1 + hs * 0.9 * math.sin(ang)
        self.line(sx, sy, ex, ey, color, sw, dash)
        if head in ("end", "both"):
            tri(x2, y2, ang)
        if head in ("start", "both"):
            tri(x1, y1, ang + math.pi)

    def polyline_arrow(self, pts, color=None, sw=0.9, hs=5.0):
        color = color or self.t["FG"]
        for (a, b), (c, d) in zip(pts[:-2], pts[1:-1]):
            self.line(a, b, c, d, color, sw)
        (a, b), (c, d) = pts[-2], pts[-1]
        self.arrow(a, b, c, d, color, sw, "end", hs)

    def label_box(self, x, y, s, size=None, bold=False, fill=None, stroke=None, r=None, pad=3.0,
                  color=None):
        """圆角框标签（think-cell CAGR / 差异箭头样式）。返回框的 (x0,y0,x1,y1)。"""
        size = size or self.t["SIZE_LABEL"]
        stroke = stroke or self.t["FG"]
        fill = fill or self.t["BG"]
        w = text_w(s, size, bold) + pad * 2 + 2
        h = size * 1.25 + pad * 1.2
        r = h / 2 if r is None else r
        self.rect(x - w / 2, y - h / 2, w, h, fill=fill, stroke=stroke, sw=0.8, r=r)
        self.text(x, y, s, size, color or self.t["FG"], "middle", "middle", bold)
        return (x - w / 2, y - h / 2, x + w / 2, y + h / 2)

    def break_mark(self, x, y, w, horizontal=True, wave=False, gap=5.0):
        """断轴标记：直线（柱形）或波浪（折线/面积/Mekko）。在 (x,y) 处横跨宽度 w。"""
        bg = self.t["BG"]
        if horizontal:
            self.rect(x - 2, y - gap / 2, w + 4, gap, fill=bg)
            for dy in (-gap / 2, gap / 2):
                if wave:
                    pts = [(x + i * w / 24, y + dy + 1.6 * math.sin(i * math.pi / 3)) for i in range(25)]
                    self.poly(pts, stroke=self.t["AXIS"], sw=0.8, closed=False)
                else:
                    self.line(x - 3, y + dy + 1.5, x + w + 3, y + dy - 1.5, self.t["AXIS"], 0.8)
        else:
            self.rect(x - gap / 2, y - 2, gap, w + 4, fill=bg)
            for dx in (-gap / 2, gap / 2):
                self.line(x + dx - 2, y - 3, x + dx + 2, y + w + 3, self.t["AXIS"], 0.8)

    # ---------------- 包围盒（碰撞检测用）
    @staticmethod
    def text_bbox(x, y, s, size, anchor="middle", valign="middle", bold=False):
        w = text_w(s, size, bold)
        h = size * 1.2
        x0 = {"start": x, "middle": x - w / 2, "end": x - w}[anchor]
        y0 = {"top": y, "middle": y - h / 2, "bottom": y - h}[valign]
        return (x0, y0, x0 + w, y0 + h)

    # ================================================================ 后端：SVG
    def to_svg(self, responsive=True) -> str:
        t = self.t
        W, H = _svg_num(self.W, "width"), _svg_num(self.H, "height")
        font, bg = _svg_font(t["FONT"]), _svg_color(t["BG"])
        out = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W:.0f} {H:.0f}" '
               + ('width="100%" ' if responsive else f'width="{W:.0f}" height="{H:.0f}" ')
               + f'font-family="{font}">']
        out.append(f'<rect x="0" y="0" width="{W:.0f}" height="{H:.0f}" fill="{bg}"/>')
        for p in self.items:
            a = p.a
            if p.kind == "rect":
                st = _svg_paint(a["fill"], a["stroke"], a["sw"], a["dash"], a["opacity"])
                x, y, w, h = (_svg_num(a[k], k) for k in ("x", "y", "w", "h"))
                r = _svg_num(a["r"], "radius")
                rr = f' rx="{r:.1f}"' if r else ""
                out.append(f'<rect x="{x:.2f}" y="{y:.2f}" width="{w:.2f}" height="{h:.2f}"{rr}{st}/>')
            elif p.kind == "line":
                x1, y1, x2, y2 = (_svg_num(a[k], k) for k in ("x1", "y1", "x2", "y2"))
                color, sw = _svg_color(a["color"]), _svg_num(a["sw"], "stroke width")
                d = f' stroke-dasharray="{_svg_dash(a["dash"])}"' if a["dash"] else ""
                out.append(f'<line x1="{x1:.2f}" y1="{y1:.2f}" x2="{x2:.2f}" y2="{y2:.2f}" '
                           f'stroke="{color}" stroke-width="{sw}"{d}/>')
            elif p.kind == "poly":
                pts = " ".join(f"{_svg_num(x, 'point x'):.2f},{_svg_num(y, 'point y'):.2f}" for x, y in a["pts"])
                tag = "polygon" if a["closed"] else "polyline"
                fill = a["fill"] if a["closed"] else None
                st = _svg_paint(fill, a["stroke"], a["sw"], a["dash"], a["opacity"])
                if not a["closed"]:
                    st += ' stroke-linejoin="round"'
                out.append(f'<{tag} points="{pts}"{st}/>')
            elif p.kind == "circle":
                st = _svg_paint(a["fill"], a["stroke"], a["sw"], None, a["opacity"])
                cx, cy, r = (_svg_num(a[k], k) for k in ("cx", "cy", "r"))
                out.append(f'<circle cx="{cx:.2f}" cy="{cy:.2f}" r="{r:.2f}"{st}/>')
            elif p.kind == "text":
                anc = a["anchor"]
                dy = {"top": "0.8em", "middle": "0.35em", "bottom": "-0.2em"}[a["valign"]]
                fw = ' font-weight="700"' if a["bold"] else ""
                it = ' font-style="italic"' if a["italic"] else ""
                x, y = _svg_num(a["x"], "text x"), _svg_num(a["y"], "text y")
                size, color = _svg_num(a["size"], "font size"), _svg_color(a["color"])
                rotate = _svg_num(a["rotate"], "rotation")
                tr = f' transform="rotate({-rotate} {x:.2f} {y:.2f})"' if rotate else ""
                out.append(f'<text x="{x:.2f}" y="{y:.2f}" dy="{dy}" text-anchor="{anc}" '
                           f'font-size="{size}" fill="{color}"{fw}{it}{tr}>{_esc(a["s"])}</text>')
        out.append("</svg>")
        return "\n".join(out)

    # ================================================================ 后端：PNG
    def to_png(self, path: str, scale: float = 2.0):
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib import font_manager
        from matplotlib.patches import Circle, FancyBboxPatch, Polygon, Rectangle

        fam = _mpl_font()
        fig = plt.figure(figsize=(self.W / 72, self.H / 72), dpi=72 * scale)
        ax = fig.add_axes([0, 0, 1, 1])
        ax.set_xlim(0, self.W)
        ax.set_ylim(self.H, 0)
        ax.axis("off")
        fig.patch.set_facecolor(self.t["BG"])
        z = 0
        for p in self.items:
            a = p.a
            z += 1
            if p.kind == "rect":
                kw = dict(facecolor=a["fill"] or "none", edgecolor=a["stroke"] or "none",
                          linewidth=a["sw"], alpha=a["opacity"], zorder=z,
                          linestyle=(0, _dash(a["dash"])) if a["dash"] else "solid")
                if a["r"]:
                    ax.add_patch(FancyBboxPatch((a["x"], a["y"]), a["w"], a["h"],
                                                boxstyle=f"round,pad=0,rounding_size={a['r']}", **kw))
                else:
                    ax.add_patch(Rectangle((a["x"], a["y"]), a["w"], a["h"], **kw))
            elif p.kind == "line":
                ax.plot([a["x1"], a["x2"]], [a["y1"], a["y2"]], color=a["color"], lw=a["sw"], zorder=z,
                        linestyle=(0, _dash(a["dash"])) if a["dash"] else "solid",
                        solid_capstyle="butt")
            elif p.kind == "poly":
                if a["closed"]:
                    ax.add_patch(Polygon(a["pts"], closed=True, facecolor=a["fill"] or "none",
                                         edgecolor=a["stroke"] or "none", linewidth=a["sw"],
                                         alpha=a["opacity"], zorder=z, joinstyle="round"))
                else:
                    xs, ys = zip(*a["pts"])
                    ax.plot(xs, ys, color=a["stroke"], lw=a["sw"], zorder=z, alpha=a["opacity"],
                            linestyle=(0, _dash(a["dash"])) if a["dash"] else "solid",
                            solid_joinstyle="round")
            elif p.kind == "circle":
                ax.add_patch(Circle((a["cx"], a["cy"]), a["r"], facecolor=a["fill"] or "none",
                                    edgecolor=a["stroke"] or "none", linewidth=a["sw"],
                                    alpha=a["opacity"], zorder=z))
            elif p.kind == "text":
                va = {"top": "top", "middle": "center", "bottom": "bottom"}[a["valign"]]
                ha = {"start": "left", "middle": "center", "end": "right"}[a["anchor"]]
                ax.text(a["x"], a["y"], a["s"], fontsize=a["size"], color=a["color"], ha=ha, va=va,
                        fontweight="bold" if a["bold"] else "normal", rotation=a["rotate"],
                        rotation_mode="anchor", family=fam, zorder=z,
                        style="italic" if a["italic"] else "normal")
        fig.savefig(path, facecolor=self.t["BG"])
        plt.close(fig)
        return path

    # ================================================================ 后端：PPTX
    def to_pptx(self, slide, x_in: float = 0.5, y_in: float = 0.5, w_in: float | None = None,
                group: bool = True):
        """把图画进 PPT：每个图元一个原生形状，整张图组合为一个组。w_in 缩放宽度。"""
        from pptx.dml.color import RGBColor
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.oxml.ns import qn
        from pptx.util import Emu, Pt

        k = (w_in * 72 / self.W) if w_in else 1.0   # pt 缩放
        ox, oy = x_in * 914400, y_in * 914400
        E = lambda v: Emu(int(round(v * k * 12700)))
        X = lambda v: Emu(int(round(ox + v * k * 12700)))
        Y = lambda v: Emu(int(round(oy + v * k * 12700)))
        grp = slide.shapes.add_group_shape() if group else None
        shapes = grp.shapes if group else slide.shapes
        t = self.t

        def rgb(h):
            h = h.lstrip("#")
            return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

        def paint(shp, fill, stroke, sw, dash=None, opacity=1.0):
            if fill:
                shp.fill.solid()
                shp.fill.fore_color.rgb = rgb(fill)
                if opacity < 1:
                    _set_alpha(shp, opacity)
            else:
                shp.fill.background()
            if stroke and sw:
                shp.line.color.rgb = rgb(stroke)
                shp.line.width = Pt(sw * k)
                if dash:
                    shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            else:
                shp.line.fill.background()
            shp.shadow.inherit = False

        def nostyle(shp):
            st = shp._element.find(qn("p:style"))
            if st is not None:
                shp._element.remove(st)
            return shp

        # 背景（保证透明区域在深色 PPT 上仍可读）
        bg = shapes.add_shape(MSO_SHAPE.RECTANGLE, X(0), Y(0), E(self.W), E(self.H))
        paint(bg, t["BG"], None, 0)
        nostyle(bg)
        for p in self.items:
            a = p.a
            if p.kind == "rect":
                if a["w"] <= 0.01 or a["h"] <= 0.01:
                    continue
                st = MSO_SHAPE.ROUNDED_RECTANGLE if a["r"] else MSO_SHAPE.RECTANGLE
                s = shapes.add_shape(st, X(a["x"]), Y(a["y"]), E(a["w"]), E(a["h"]))
                if a["r"]:
                    s.adjustments[0] = min(0.5, a["r"] / max(min(a["w"], a["h"]), 0.01))
                paint(s, a["fill"], a["stroke"], a["sw"], a["dash"], a["opacity"])
                nostyle(s)
            elif p.kind == "line":
                c = nostyle(shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X(a["x1"]), Y(a["y1"]), X(a["x2"]), Y(a["y2"])))
                c.line.color.rgb = rgb(a["color"])
                c.line.width = Pt(a["sw"] * k)
                if a["dash"]:
                    c.line.dash_style = MSO_LINE_DASH_STYLE.SQUARE_DOT if _dash(a["dash"])[0] <= 2 \
                        else MSO_LINE_DASH_STYLE.DASH
            elif p.kind == "poly":
                pts = a["pts"]
                fb = shapes.build_freeform(X(pts[0][0]), Y(pts[0][1]), scale=1.0)
                fb.add_line_segments([(X(x), Y(y)) for x, y in pts[1:]], close=a["closed"])
                s = nostyle(fb.convert_to_shape())
                paint(s, a["fill"] if a["closed"] else None, a["stroke"], a["sw"], a["dash"], a["opacity"])
            elif p.kind == "circle":
                r = a["r"]
                s = nostyle(shapes.add_shape(MSO_SHAPE.OVAL, X(a["cx"] - r), Y(a["cy"] - r), E(2 * r), E(2 * r)))
                paint(s, a["fill"], a["stroke"], a["sw"], None, a["opacity"])
            elif p.kind == "text":
                size = a["size"]
                w = text_w(a["s"], size, a["bold"]) * 1.08 + 4
                h = size * 1.45
                x0 = {"start": a["x"] - 2, "middle": a["x"] - w / 2, "end": a["x"] - w + 2}[a["anchor"]]
                y0 = {"top": a["y"] - size * 0.12, "middle": a["y"] - h / 2, "bottom": a["y"] - h + size * 0.12}[a["valign"]]
                if a["rotate"]:
                    # 旋转绕中心：先把框中心放在锚点处
                    x0, y0 = a["x"] - w / 2, a["y"] - h / 2
                tb = shapes.add_textbox(X(x0), Y(y0), E(w), E(h))
                tf = tb.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = False
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                para = tf.paragraphs[0]
                para.alignment = {"start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}[a["anchor"]]
                run = para.add_run()
                run.text = a["s"]
                f = run.font
                f.size = Pt(size * k)
                f.bold = a["bold"]
                f.italic = a["italic"]
                f.color.rgb = rgb(a["color"])
                f.name = t["FONT_PPT_LATIN"]
                rPr = run._r.get_or_add_rPr()
                ea = rPr.find(qn("a:ea"))
                if ea is None:
                    ea = rPr.makeelement(qn("a:ea"), {})
                    rPr.append(ea)
                ea.set("typeface", t["FONT_PPT_EA"])
                if a["rotate"]:
                    tb.rotation = -a["rotate"]
        return grp or shapes


# ---------------------------------------------------------------- 工具
def _esc(s: str) -> str:
    return html.escape(str(s), quote=False)


def _svg_num(value, label):
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"invalid SVG {label}") from exc
    if not math.isfinite(number):
        raise ValueError(f"invalid SVG {label}")
    return number


def _svg_color(value):
    color = str(value)
    if not re.fullmatch(r"#[0-9A-Fa-f]{6}", color):
        raise ValueError("invalid SVG color")
    return color


def _svg_font(value):
    font = str(value)
    if len(font) > 500 or re.search(r"[<>&;`]", font) or re.search(r"(?:onload|onerror|javascript)\s*[:=]", font, re.I):
        raise ValueError("invalid SVG font")
    return html.escape(font, quote=True)


def _svg_dash(value):
    try:
        parts = [float(v) for v in str(value).replace(",", " ").split()]
    except ValueError as exc:
        raise ValueError("invalid SVG dash") from exc
    if not parts or any(not math.isfinite(v) or v < 0 for v in parts):
        raise ValueError("invalid SVG dash")
    return " ".join(f"{v:g}" for v in parts)


def _dash(d):
    if isinstance(d, str):
        return tuple(float(v) for v in d.replace(",", " ").split())
    return tuple(d)


def _svg_paint(fill, stroke, sw, dash, opacity):
    s = f' fill="{_svg_color(fill)}"' if fill else ' fill="none"'
    if stroke and sw:
        s += f' stroke="{_svg_color(stroke)}" stroke-width="{_svg_num(sw, "stroke width")}"'
        if dash:
            s += f' stroke-dasharray="{_svg_dash(dash)}"'
    opacity = _svg_num(opacity, "opacity") if opacity is not None else None
    if opacity is not None and not 0 <= opacity <= 1:
        raise ValueError("invalid SVG opacity")
    if opacity is not None and opacity < 1:
        s += f' opacity="{opacity}"'
    return s


def _set_alpha(shp, opacity):
    from pptx.oxml.ns import qn
    sf = shp.fill._xPr.find(qn("a:solidFill"))
    if sf is None:
        return
    clr = sf[0]
    al = clr.makeelement(qn("a:alpha"), {"val": str(int(opacity * 100000))})
    clr.append(al)


_FONT_CACHE = None


def _mpl_font():
    global _FONT_CACHE
    if _FONT_CACHE:
        return _FONT_CACHE
    import platform
    from matplotlib import font_manager
    prefs = {"Darwin": ["PingFang SC", "Hiragino Sans GB", "Heiti SC"],
             "Windows": ["Microsoft YaHei", "SimHei"]}.get(platform.system(), [])
    prefs += ["Noto Sans CJK SC", "Source Han Sans SC", "WenQuanYi Micro Hei", "Noto Sans CJK JP"]
    have = {f.name for f in font_manager.fontManager.ttflist}
    fam = [p for p in prefs if p in have][:1] + ["DejaVu Sans"]
    import matplotlib
    matplotlib.rcParams["axes.unicode_minus"] = False
    _FONT_CACHE = fam
    return fam
