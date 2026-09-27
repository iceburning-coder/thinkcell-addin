"""插件桥接层：JS 传入图表 JSON（与报告 JSON 中单张图的写法相同），返回 SVG / PPTX(base64)。"""
import base64
import io
import json

from tccore.report import build_figure


def _fig(spec_json):
    ch = json.loads(spec_json)
    return build_figure(ch, ch.get("theme"))


def render(spec_json):
    """→ JSON {svg, width, height, meta}"""
    f = _fig(spec_json)
    return json.dumps({"svg": f.c.to_svg(False), "width": f.W, "height": f.H,
                       "meta": [{"annotation": m["annotation"], "value": m["value"]}
                                for m in f.meta if isinstance(m["value"], (int, float, str))]},
                      ensure_ascii=False)


def render_pptx(spec_json, slide_w_in=13.333, slide_h_in=7.5):
    """→ base64 编码的单页 pptx（原生形状，可直接编辑），用于 insertSlidesFromBase64。"""
    from pptx import Presentation
    from pptx.util import Inches
    f = _fig(spec_json)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(slide_w_in), Inches(slide_h_in)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    margin = 0.35
    w = min(slide_w_in - 2 * margin, (slide_h_in - 2 * margin) * f.W / f.H)
    h = w * f.H / f.W
    f.pptx(s, (slide_w_in - w) / 2, (slide_h_in - h) / 2, w_in=w)
    bio = io.BytesIO()
    prs.save(bio)
    return base64.b64encode(bio.getvalue()).decode()
